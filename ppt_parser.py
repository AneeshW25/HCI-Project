from pptx import Presentation

def extract_slides(file_path):
    prs = Presentation(file_path)
    slides_data = []

    for i, slide in enumerate(prs.slides):
        title = ""
        content = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if not title:
                    title = text
                else:
                    content.append(text)

        slides_data.append({
            "slide_number": i + 1,
            "title": title,
            "content": " ".join(content)
        })

    return slides_data